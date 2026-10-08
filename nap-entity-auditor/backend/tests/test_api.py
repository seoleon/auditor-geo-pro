"""Pruebas de integración de la API: alta, auditoría completa, aislamiento, exportaciones e historial."""
import io
import json

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from tests.conftest import FixtureFetcher, fixture_json

from app.core.config import get_settings
from app.main import app
from app.services.discovery.providers import BraveSearchProvider
from app.services.integrations.google import GooglePlacesClient

H = {"X-Requested-With": "nap-auditor"}

BUSINESS = {
    "official_name": "Centro Aurora Bienestar",
    "name_variants": ["Aurora Bienestar"],
    "domain": "https://www.aurora-bienestar.es/",
    "country": "ES",
    "city": "Valencia",
    "primary_category": "Spa",
    "business_type": "physical",
    "nap_name": "Centro Aurora Bienestar",
    "address_street": "Calle de Colón, 12, 3º, pta 4",
    "postal_code": "46004",
    "locality": "Valencia",
    "nap_province": "Valencia",
    "phone_primary": "627 171 728",
    "old_phones": ["963 11 22 33"],
    "website": "https://www.aurora-bienestar.es/",
    "opening_hours": "Mo-Fr 10:00-20:00; Sa 10:00-14:00",
    "social_profiles": {"facebook": "https://www.facebook.com/aurorabienestar", "instagram": "https://www.instagram.com/aurorabienestar/"},
}


def client_for(email: str, org: str = "Agencia") -> TestClient:
    c = TestClient(app)
    r = c.post("/api/auth/register", json={"email": email, "password": "contraseña-segura-123", "organization": org})
    assert r.status_code == 200, r.text
    return c


@pytest.fixture
def mocked_world(monkeypatch):
    """Modo real con proveedores simulados: sin red ni llamadas de pago."""
    s = get_settings().model_copy(update={"BRAVE_API_KEY": "test", "GOOGLE_PLACES_API_KEY": "test", "SEARCH_MAX_PAGES_PER_QUERY": 1})
    brave = BraveSearchProvider(s, transport=httpx.MockTransport(lambda r: httpx.Response(200, json=json.loads(fixture_json("brave_page1.json")))))
    places_t = httpx.MockTransport(lambda r: httpx.Response(200, json=json.loads(fixture_json("places_text.json"))))
    fetcher = FixtureFetcher()
    monkeypatch.setattr("app.services.audit_runner.configured_providers", lambda settings=None: [brave])
    monkeypatch.setattr("app.services.audit_runner.SafeFetcher", lambda: fetcher)
    monkeypatch.setattr("app.services.audit_runner.GooglePlacesClient", lambda settings=None: GooglePlacesClient(s, transport=places_t))
    return fetcher


def create_confirmed_business(c: TestClient) -> dict:
    r = c.post("/api/businesses", json=BUSINESS, headers=H)
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["domain"] == "aurora-bienestar.es" and b["nap_confirmed"] is False
    r = c.post(f"/api/businesses/{b['id']}/confirm-nap", headers=H)
    assert r.status_code == 200 and r.json()["nap_confirmed"] is True
    return r.json()


def test_auth_required_and_csrf():
    c = TestClient(app)
    assert c.get("/api/businesses").status_code == 401
    c = client_for("csrf@test.es")
    assert c.post("/api/clients", json={"name": "X"}).status_code == 403  # sin cabecera anti-CSRF
    assert c.post("/api/clients", json={"name": "X"}, headers=H).status_code == 200


def test_login_and_wrong_password():
    client_for("login@test.es")
    c = TestClient(app)
    assert c.post("/api/auth/login", json={"email": "login@test.es", "password": "mal"}).status_code == 401
    r = c.post("/api/auth/login", json={"email": "LOGIN@test.es", "password": "contraseña-segura-123"})
    assert r.status_code == 200 and c.get("/api/auth/me").json()["email"] == "login@test.es"


def test_business_validation_errors():
    c = client_for("val@test.es")
    bad = dict(BUSINESS, phone_primary="123", postal_code="4600", domain="no es un dominio")
    r = c.post("/api/businesses", json=bad, headers=H)
    assert r.status_code == 422
    fields = {e["field"] for e in r.json()["errors"]}
    assert "domain" in fields


def test_confirm_nap_requires_complete_data():
    c = client_for("conf@test.es")
    b = c.post("/api/businesses", json={"official_name": "Sadhana Center", "domain": "sadhanacenter.com"}, headers=H).json()
    r = c.post(f"/api/businesses/{b['id']}/confirm-nap", headers=H)
    assert r.status_code == 422 and "teléfono" in r.json()["detail"]


def test_tenant_isolation():
    a = client_for("a@test.es", "Agencia A")
    b = client_for("b@test.es", "Agencia B")
    biz = create_confirmed_business(a)
    assert b.get(f"/api/businesses/{biz['id']}").status_code == 404
    assert b.put(f"/api/businesses/{biz['id']}", json=BUSINESS, headers=H).status_code == 404
    assert b.post(f"/api/businesses/{biz['id']}/audits", json={"mode": "demo"}, headers=H).status_code == 404
    assert b.get("/api/businesses").json() == []
    r = a.post(f"/api/businesses/{biz['id']}/audits", json={"mode": "demo"}, headers=H)
    audit_id = r.json()["id"]
    assert b.get(f"/api/audits/{audit_id}").status_code == 404
    assert b.get(f"/api/audits/{audit_id}/sources").status_code == 404
    assert b.get(f"/api/audits/{audit_id}/export.csv").status_code == 404
    src = a.get(f"/api/audits/{audit_id}/sources").json()["items"][0]
    assert b.get(f"/api/sources/{src['id']}").status_code == 404
    assert b.patch(f"/api/sources/{src['id']}", json={"manual_status": "CORRECT"}, headers=H).status_code == 404
    # Un cliente de otra organización tampoco puede asignarse
    other_client = b.post("/api/clients", json={"name": "Cliente B"}, headers=H).json()
    assert a.post("/api/businesses", json=dict(BUSINESS, client_id=other_client["id"]), headers=H).status_code == 404


def test_demo_audit_flow_and_exports():
    c = client_for("demo@test.es")
    biz = create_confirmed_business(c)
    r = c.post(f"/api/businesses/{biz['id']}/audits", json={"mode": "demo"}, headers=H)
    assert r.status_code == 200
    audit = c.get(f"/api/audits/{r.json()['id']}").json()
    assert audit["status"] in ("completed", "partial"), audit
    assert any("SIMULADOS" in lim for lim in audit["limitations"])
    sm = audit["summary"]
    assert sm["demo"] is True and sm["confirmed_inconsistencies"] >= 1 and sm["possible_duplicate_groups"] >= 1
    items = c.get(f"/api/audits/{audit['id']}/sources?page_size=200").json()["items"]
    assert all(s["extracted"].get("demo") for s in items if s["fetch_status"] in ("ok", "api"))
    csv = c.get(f"/api/audits/{audit['id']}/export.csv")
    assert csv.status_code == 200 and "DATOS SIMULADOS" in csv.content.decode("utf-8-sig").splitlines()[0]
    assert "DEMO" in csv.headers["content-disposition"]


def test_real_mode_audit_with_mocked_providers(mocked_world):
    c = client_for("real@test.es")
    biz = create_confirmed_business(c)
    r = c.post(f"/api/businesses/{biz['id']}/audits", json={"mode": "real"}, headers=H)
    audit = c.get(f"/api/audits/{r.json()['id']}").json()
    assert audit["status"] in ("completed", "partial"), audit
    sm = audit["summary"]
    srcs = c.get(f"/api/audits/{audit['id']}/sources?page_size=200").json()["items"]
    by_url = {s["url"]: s for s in srcs}

    # Web oficial: la página de contacto con teléfono antiguo es P0
    contact = by_url["https://www.aurora-bienestar.es/contacto/"]
    assert contact["overall_status"] == "CONFIRMED_INCONSISTENCY" and contact["priority"] == "P0"
    # Directorio con teléfono antiguo: inconsistencia confirmada
    old = by_url["https://empresas-valencia.com/empresa/centro-aurora-bienestar"]
    assert old["overall_status"] == "CONFIRMED_INCONSISTENCY"
    # Variante tipográfica de dirección: no crítica
    guia = next(s for s in srcs if "guiacomercial.es" in s["url"])
    assert guia["priority"] in ("P3", None)
    # Bloqueo: no verificable
    assert by_url["https://bloqueada.es/ficha/aurora"]["overall_status"] == "UNVERIFIABLE"
    # Irrelevante
    assert by_url["https://talleresperez.es/"]["source_type"] == "irrelevant"
    # Perfiles sociales aportados por el usuario: robots.txt los bloquea -> no verificables
    assert by_url["https://www.facebook.com/aurorabienestar"]["overall_status"] == "UNVERIFIABLE"
    # Places API: ficha principal y duplicado
    assert audit["gbp_report"]["main_listing"]["place_id"] == "ChIJ_TEST_MAIN"
    dups = c.get(f"/api/audits/{audit['id']}/duplicates").json()
    platforms = {d["platform"] for d in dups}
    assert "Google Maps" in platforms and "guialocal.es" in platforms
    # Datos estructurados: la página de servicios contradice el NAP
    codes = {i["code"] for i in audit["schema_report"]["issues"]}
    assert "old_phone" in codes and "address_mismatch" in codes
    # Grafo con evidencias
    g = audit["entity_graph"]
    assert g["nodes"] and all(e["evidence"] for e in g["edges"])
    # GEO
    assert audit["geo_report"]["checks"] and "No garantizan" in audit["geo_report"]["disclaimer"]
    assert audit["geo_report"]["robots_ai"]["bots"]["GPTBot"] is False
    # Acciones priorizadas con certeza
    actions = c.get(f"/api/audits/{audit['id']}/actions").json()
    assert actions[0]["priority"] == "P0"
    assert {"confirmed", "hypothesis"} <= {a["certainty"] for a in actions}
    assert sm["queries"]["total"] > 0 and sm["verified_citations"] >= 1
    # Evidencia disponible en el detalle de la fuente
    det = c.get(f"/api/sources/{old['id']}").json()
    assert det["evidence"]["phone"] and det["field_notes"]["phone"] and det["discovered_by"][0]["provider"] == "brave"

    # Filtros, búsqueda, ordenación y paginación
    page = c.get(f"/api/audits/{audit['id']}/sources?status=CONFIRMED_INCONSISTENCY&page_size=1").json()
    assert page["total"] >= 2 and len(page["items"]) == 1
    q = c.get(f"/api/audits/{audit['id']}/sources?q=guialocal").json()
    assert q["total"] == 2
    srt = c.get(f"/api/audits/{audit['id']}/sources?sort=confidence&order=desc&page_size=200").json()["items"]
    assert [s["confidence"] for s in srt] == sorted([s["confidence"] for s in srt], reverse=True)

    # Exportaciones
    xlsx = c.get(f"/api/audits/{audit['id']}/export.xlsx")
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Resumen ejecutivo", "NAP oficial", "Citaciones encontradas", "Inconsistencias", "Posibles duplicados",
                             "Datos estructurados", "Perfiles sociales", "Acciones recomendadas", "Fuentes no verificables"]
    pdf = c.get(f"/api/audits/{audit['id']}/export.pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF") and len(pdf.content) > 3000
    csv = c.get(f"/api/audits/{audit['id']}/export.csv").content.decode("utf-8-sig")
    assert "Evidencias" in csv.splitlines()[0] and "empresas-valencia.com" in csv

    # Decisión manual sobre duplicados y acciones
    d = dups[0]
    assert c.patch(f"/api/duplicates/{d['id']}", json={"status": "dismissed", "note": "Son ubicaciones distintas"}, headers=H).json()["status"] == "dismissed"
    assert c.patch(f"/api/actions/{actions[0]['id']}", json={"status": "done"}, headers=H).json()["status"] == "done"
    assert c.patch(f"/api/sources/{guia['id']}", json={"manual_status": "CORRECT", "manual_note": "Revisado"}, headers=H).json()["overall_status"] == "CORRECT"

    # Segunda auditoría: historial, comparación y herencia de decisiones
    r2 = c.post(f"/api/businesses/{biz['id']}/audits", json={"mode": "economic"}, headers=H)
    a2 = c.get(f"/api/audits/{r2.json()['id']}").json()
    assert a2["status"] in ("completed", "partial")
    cmp = c.get(f"/api/audits/compare?old={audit['id']}&new={a2['id']}").json()
    assert "no demuestra que hayan desaparecido" in cmp["not_observed_note"]
    dups2 = c.get(f"/api/audits/{a2['id']}/duplicates").json()
    same = [x for x in dups2 if x["platform"] == d["platform"]]
    assert same and same[0]["status"] == "dismissed"  # la decisión manual se conserva
    reused = [s for s in c.get(f"/api/audits/{a2['id']}/sources?page_size=200").json()["items"] if (s["render_method"] or "").startswith("reused")]
    assert reused  # modo económico reutiliza resultados recientes

    dash = c.get("/api/dashboard").json()
    assert dash["businesses"] == 1 and dash["audits"] == 2 and dash["history"]


def test_unconfirmed_nap_audit_does_not_compare(mocked_world):
    c = client_for("unconf@test.es")
    b = c.post("/api/businesses", json=dict(BUSINESS), headers=H).json()
    r = c.post(f"/api/businesses/{b['id']}/audits", json={"mode": "real"}, headers=H)
    audit = c.get(f"/api/audits/{r.json()['id']}").json()
    assert audit["summary"]["compared"] is False
    assert any("no está confirmado" in x for x in audit["limitations"])
    srcs = c.get(f"/api/audits/{audit['id']}/sources?page_size=200").json()["items"]
    assert not any(s["overall_status"] == "CONFIRMED_INCONSISTENCY" for s in srcs)
    assert c.get(f"/api/audits/{audit['id']}/actions").json() == []
    detected = c.get(f"/api/businesses/{b['id']}/detected-nap").json()
    assert detected["phones"] and "NO se asumen correctos" in detected["note"]


def test_no_providers_configured_is_documented(monkeypatch):
    monkeypatch.setattr("app.services.audit_runner.SafeFetcher", lambda: FixtureFetcher())
    c = client_for("noprov@test.es")
    biz = create_confirmed_business(c)
    r = c.post(f"/api/businesses/{biz['id']}/audits", json={"mode": "real"}, headers=H)
    audit = c.get(f"/api/audits/{r.json()['id']}").json()
    lims = " ".join(audit["limitations"])
    assert "proveedor de búsqueda" in lims and "GOOGLE_PLACES_API_KEY" in lims
    assert audit["gbp_report"]["authorized"] is False
    assert audit["summary"]["sources_fetched"] >= 1  # la web oficial sí se audita


def test_nap_change_requires_reconfirmation_and_history():
    c = client_for("hist@test.es")
    biz = create_confirmed_business(c)
    upd = dict(BUSINESS, phone_primary="611 222 333")
    r = c.put(f"/api/businesses/{biz['id']}", json=upd, headers=H).json()
    assert r["nap_confirmed"] is False
    changes = c.get(f"/api/businesses/{biz['id']}/changes").json()
    assert any(ch["field"] == "phone_primary" for ch in changes)
    r = c.post(f"/api/businesses/{biz['id']}/variants", json={"variant": "Aurora Spa", "decision": "reject"}, headers=H).json()
    assert r["rejected_name_variants"] == ["Aurora Spa"]


def test_ai_test_registration_detects_errors():
    c = client_for("ai@test.es")
    biz = create_confirmed_business(c)
    resp = ("Centro Aurora Bienestar es un spa en Valencia. Puedes llamar al 963 11 22 33. Está en Calle de Colón 14, 46004 Valencia. "
            "Más info en https://www.aurora-bienestar.es/")
    r = c.post(f"/api/businesses/{biz['id']}/ai-tests", json={"provider": "perplexity", "model": "sonar", "query": "spa en Valencia",
                                                               "response": resp}, headers=H)
    t = r.json()
    assert t["mentions"]["name"] is True and t["mentions"]["domain"] is True
    assert any("antiguo" in e for e in t["errors_detected"])
    assert any("Dirección distinta" in e for e in t["errors_detected"])
    run = c.post(f"/api/businesses/{biz['id']}/ai-tests/run", json={"provider": "chatgpt", "query": "x?"}, headers=H)
    assert run.status_code == 400 and "no configurada" in run.json()["detail"]


def test_settings_never_expose_keys(monkeypatch):
    c = client_for("keys@test.es")
    r = c.get("/api/settings/providers").text
    assert "test-secret-key" not in r and "api_key" not in r.lower()


def test_scheduler_creates_due_audits(db):
    from app.models import Business, Organization
    from app.tasks import run_due_audits

    org = Organization(name="o")
    db.add(org)
    db.flush()
    db.add(Business(organization_id=org.id, official_name="X", domain="x.es", audit_frequency="weekly", nap_confirmed=True,
                    audit_mode_default="demo", phone_primary="627171728"))
    db.add(Business(organization_id=org.id, official_name="Y", domain="y.es", audit_frequency="weekly", nap_confirmed=False))
    db.commit()
    created = run_due_audits()
    assert len(created) == 1
    assert run_due_audits() == []  # la siguiente ejecución queda programada a 7 días
