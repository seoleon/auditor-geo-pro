"""Exportación de informes: CSV, Excel (9 hojas) y PDF."""
from __future__ import annotations

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import SOURCE_TYPE_LABELS_ES, STATUS_LABELS_ES
from app.services.reports.data import SOURCE_COLUMNS, ReportData, source_row

DEMO_BANNER = "DATOS SIMULADOS — MODO DEMO: este informe no describe datos reales"
HEADER_FILL = PatternFill("solid", fgColor="1E293B")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _safe_cell(v):
    """Evita inyección de fórmulas en CSV/Excel (valores que empiezan por = + - @)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def export_csv(data: ReportData) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    if data.demo:
        w.writerow([DEMO_BANNER])
    w.writerow([label for _, label in SOURCE_COLUMNS])
    for s in data.sources:
        row = source_row(s)
        w.writerow([_safe_cell(row[k]) for k, _ in SOURCE_COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], demo: bool, widths: dict[int, int] | None = None):
    ws = wb.create_sheet(title[:31])
    r0 = 1
    if demo:
        ws.cell(row=1, column=1, value=DEMO_BANNER).font = Font(color="B91C1C", bold=True)
        r0 = 2
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=r0, column=i, value=h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for ri, row in enumerate(rows, r0 + 1):
        for ci, v in enumerate(row, 1):
            if isinstance(v, (list, dict)):
                v = str(v)
            cell = ws.cell(row=ri, column=ci, value=_safe_cell(v))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 22)
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
    return ws


def export_excel(data: ReportData) -> bytes:
    a, b, sm = data.audit, data.business, data.audit.summary or {}
    wb = Workbook()
    wb.remove(wb.active)
    # 1. Resumen ejecutivo
    rows = [
        ["Empresa", b.official_name], ["Dominio", b.domain], ["Auditoría", f"#{a.id}"], ["Modo", a.mode],
        ["Fecha", (a.finished_at or a.created_at).strftime("%Y-%m-%d %H:%M")], ["Estado", a.status],
        ["NAP oficial confirmado", "Sí" if sm.get("compared") else "No (no se comparan discrepancias)"],
        ["Fuentes descubiertas", sm.get("sources_total")], ["Fuentes consultadas", sm.get("sources_fetched")],
        ["Fuentes atribuidas al negocio", sm.get("attributed_sources")], ["Citaciones verificadas", sm.get("verified_citations")],
        ["Inconsistencias confirmadas", sm.get("confirmed_inconsistencies")], ["Posibles inconsistencias", sm.get("possible_inconsistencies")],
        ["Grupos de posibles duplicados", sm.get("possible_duplicate_groups")], ["Pendientes de revisión", sm.get("pending_review")],
        ["No verificables", sm.get("sources_unverifiable")], ["Consistencia NAP (interna, %)", sm.get("nap_consistency_pct")],
        ["Puntuación de señales GEO (interna)", sm.get("geo_signal_score")],
        ["Aviso", "Las puntuaciones son internas y basadas en evidencias; no son factores de ranking de Google."],
    ]
    for lim in a.limitations or []:
        rows.append(["Limitación", lim])
    _sheet(wb, "Resumen ejecutivo", ["Indicador", "Valor"], rows, data.demo, {1: 38, 2: 90})
    # 2. NAP oficial
    snap = a.nap_snapshot or {}
    nap_rows = [[k, ", ".join(map(str, v)) if isinstance(v, list) else (str(v) if isinstance(v, dict) else v)] for k, v in snap.items()]
    _sheet(wb, "NAP oficial", ["Campo", "Valor"], nap_rows, data.demo, {1: 28, 2: 80})
    # 3. Citaciones
    headers = [label for _, label in SOURCE_COLUMNS]
    all_rows = [[source_row(s)[k] for k, _ in SOURCE_COLUMNS] for s in data.sources if s.fetch_status in ("ok", "api")]
    _sheet(wb, "Citaciones encontradas", headers, all_rows, data.demo, {2: 50, 18: 60, 21: 80})
    # 4. Inconsistencias
    inc = [s for s in data.sources if s.overall_status in ("CONFIRMED_INCONSISTENCY", "POSSIBLE_INCONSISTENCY")]
    _sheet(wb, "Inconsistencias", ["Certeza", "Fuente", "URL", "Campos", "Detalle", "Evidencia", "Prioridad"], [
        ["CONFIRMADA" if s.overall_status == "CONFIRMED_INCONSISTENCY" else "HIPÓTESIS (posible)", source_row(s)["source"], s.url,
         ", ".join(f"{k}: {STATUS_LABELS_ES.get(v, v)}" for k, v in (s.field_status or {}).items() if "INCONSISTENCY" in v),
         " | ".join(f"{k}: {'; '.join(v)}" for k, v in (s.field_notes or {}).items() if v),
         " | ".join(f"{k}: {v}" for k, v in (s.evidence or {}).items() if v), s.priority]
        for s in inc], data.demo, {3: 50, 5: 70, 6: 70})
    # 5. Duplicados
    _sheet(wb, "Posibles duplicados", ["Grupo", "Plataforma", "Estado", "URLs", "Motivos", "Advertencias", "Prioridad"], [
        [g.id, g.platform, g.status, "\n".join(m["url"] for m in g.members), "; ".join(g.reasons), "; ".join(g.warnings), g.priority]
        for g in data.groups], data.demo, {4: 60, 5: 60, 6: 60})
    # 6. Datos estructurados
    sr = a.schema_report or {}
    rows6 = [["Entidad", e.get("page_url"), ", ".join(e.get("types") or []), str(e.get("properties"))[:3000], ""] for e in sr.get("entities") or []]
    rows6 += [[f"Incidencia ({i['severity']})", i.get("page_url"), i.get("code"), i.get("message"), i.get("evidence") or ""] for i in sr.get("issues") or []]
    rows6 += [["Recomendación", "", "", r, ""] for r in sr.get("recommendations") or []]
    _sheet(wb, "Datos estructurados", ["Tipo de fila", "Página", "Tipo / código", "Detalle", "Evidencia"], rows6, data.demo, {2: 45, 4: 90, 5: 50})
    # 7. Perfiles sociales
    soc = (a.social_report or {}).get("profiles") or []
    _sheet(wb, "Perfiles sociales", ["Plataforma", "URL", "Estado", "Registrado", "Enlazado desde web", "En sameAs", "Estado NAP", "Consulta"], [
        [p.get("platform"), p.get("url"), p.get("status"), "Sí" if p.get("known") else "No", "Sí" if p.get("linked_from_website") else "No",
         "Sí" if p.get("in_same_as") else "No", STATUS_LABELS_ES.get(p.get("nap_status"), p.get("nap_status")), p.get("fetch_status")]
        for p in soc], data.demo, {2: 55})
    # 8. Acciones
    _sheet(wb, "Acciones recomendadas", ["Prioridad", "Categoría", "Certeza", "Acción", "Detalle", "URL", "Estado"], [
        [x.priority, x.category, "Confirmada" if x.certainty == "confirmed" else "Hipótesis", x.title, x.detail, x.url, x.status]
        for x in data.actions], data.demo, {4: 60, 5: 80, 6: 50})
    # 9. No verificables
    unv = [s for s in data.sources if s.fetch_status not in ("ok", "api")]
    _sheet(wb, "Fuentes no verificables", ["Fuente", "URL", "Tipo", "Motivo", "HTTP"], [
        [source_row(s)["source"], s.url, SOURCE_TYPE_LABELS_ES.get(s.source_type, s.source_type), f"{s.fetch_status}: {s.fetch_detail or ''}", s.http_status]
        for s in unv], data.demo, {2: 60, 4: 80})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_pdf(data: ReportData) -> bytes:
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.charts.piecharts import Pie
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    a, b, sm = data.audit, data.business, data.audit.summary or {}
    styles = getSampleStyleSheet()
    h1, h2, body = styles["Title"], styles["Heading2"], styles["BodyText"]
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10)
    center = ParagraphStyle("center", parent=body, alignment=TA_CENTER, fontSize=12)
    warn = ParagraphStyle("warn", parent=body, textColor=colors.HexColor("#B91C1C"), fontSize=11)

    def P(text, st=body):
        from xml.sax.saxutils import escape

        return Paragraph(escape(str(text or "")), st)

    def table(rows, widths, header=True):
        t = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
        style = [("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                 ("FONTSIZE", (0, 0), (-1, -1), 8)]
        if header:
            style += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white)]
        t.setStyle(TableStyle(style))
        return t

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(1.5 * cm, 1 * cm, f"NAP Entity Auditor Pro — {b.official_name} — auditoría #{a.id}")
        canvas.drawRightString(19.5 * cm, 1 * cm, f"Página {doc.page}")
        if data.demo:
            canvas.setFillColor(colors.HexColor("#B91C1C"))
            canvas.drawCentredString(10.5 * cm, 28.7 * cm, DEMO_BANNER)
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.8 * cm, bottomMargin=1.6 * cm,
                            title=f"Auditoría NAP — {b.official_name}")
    story = []
    date = (a.finished_at or a.created_at).strftime("%d/%m/%Y")
    # Portada
    story += [Spacer(1, 5 * cm), P("Auditoría de identidad digital y consistencia NAP", h1), Spacer(1, 1 * cm),
              P(b.official_name, center), P(b.domain, center), P(f"Fecha de auditoría: {date}", center), Spacer(1, 1 * cm)]
    if data.demo:
        story.append(P(DEMO_BANNER, warn))
    if not sm.get("compared"):
        story.append(P("El NAP oficial no estaba confirmado: este informe recoge fuentes y datos detectados, sin comparación.", warn))
    story.append(PageBreak())
    # Resumen ejecutivo
    story.append(P("1. Resumen ejecutivo", h2))
    kpis = [["Indicador", "Valor"],
            ["Fuentes descubiertas", sm.get("sources_total")], ["Fuentes atribuidas al negocio", sm.get("attributed_sources")],
            ["Citaciones verificadas (coherentes)", sm.get("verified_citations")],
            ["Inconsistencias confirmadas", sm.get("confirmed_inconsistencies")], ["Posibles inconsistencias (hipótesis)", sm.get("possible_inconsistencies")],
            ["Grupos de posibles duplicados", sm.get("possible_duplicate_groups")], ["Fuentes no verificables", sm.get("sources_unverifiable")],
            ["Consistencia NAP (puntuación interna)", f"{sm.get('nap_consistency_pct')} %" if sm.get("nap_consistency_pct") is not None else "n/d"]]
    story.append(table([[P(x, small) for x in r] for r in kpis], [10 * cm, 6 * cm]))
    story.append(Spacer(1, 0.4 * cm))
    p0 = [x for x in data.actions if x.priority == "P0"]
    story.append(P(f"Acciones críticas (P0): {len(p0)}. Acciones totales: {len(data.actions)}.", body))
    # Gráficos
    by_status = {k: v for k, v in (sm.get("by_status") or {}).items() if v}
    if by_status:
        d = Drawing(17 * cm, 6.5 * cm)
        pie = Pie()
        pie.x, pie.y, pie.width, pie.height = 30, 20, 140, 140
        pie.data = list(by_status.values())
        pie.labels = [f"{STATUS_LABELS_ES.get(k, k)} ({v})" for k, v in by_status.items()]
        pie.sideLabels = True
        pie.slices.fontSize = 6
        palette = ["#16A34A", "#65A30D", "#DC2626", "#F59E0B", "#7C3AED", "#64748B", "#0EA5E9", "#94A3B8", "#EA580C", "#CBD5E1"]
        for i in range(len(pie.data)):
            pie.slices[i].fillColor = colors.HexColor(palette[i % len(palette)])
        d.add(pie)
        d.add(String(300, 170, "Estado NAP de las fuentes", fontSize=9))
        bp = sm.get("actions_by_priority") or {}
        bc = VerticalBarChart()
        bc.x, bc.y, bc.width, bc.height = 300, 30, 160, 120
        bc.data = [[bp.get(p, 0) for p in ("P0", "P1", "P2", "P3")]]
        bc.categoryAxis.categoryNames = ["P0", "P1", "P2", "P3"]
        bc.valueAxis.valueMin = 0
        bc.bars[0].fillColor = colors.HexColor("#1E293B")
        d.add(bc)
        story.append(d)
    # Metodología
    story.append(P("2. Metodología", h2))
    for t in [
        "Se partió de los datos oficiales (NAP) confirmados por el usuario. La web oficial se analizó página a página (datos estructurados, HTML semántico y texto visible).",
        "Las menciones se descubrieron con proveedores de búsqueda configurados mediante combinaciones de nombre, ciudad, dirección, teléfono, dominio y variantes. "
        "Los resultados son una muestra: que una fuente no aparezca no implica que no exista.",
        "Cada página pública accesible se consultó respetando robots.txt; las páginas bloqueadas, con CAPTCHA o que exigen inicio de sesión se clasifican como NO VERIFICABLES, sin intentar eludir protecciones.",
        "Los datos públicos de Google se obtienen solo mediante Google Places API; los datos privados de Business Profile solo con autorización OAuth.",
        "Nombre, dirección, teléfono, web y horario se comparan por separado. Una discrepancia es CONFIRMADA solo si la fuente está atribuida al negocio con evidencias suficientes; en otro caso es una HIPÓTESIS.",
        "La confianza y la consistencia son puntuaciones internas basadas en evidencias, no factores de ranking de Google.",
    ]:
        story.append(P("• " + t, body))
    # Resultados
    story.append(P("3. Resultados por fuente", h2))
    rows = [["Fuente", "Tipo", "Estado", "Conf.", "Prior."]]
    for s in data.sources[:150]:
        r = source_row(s)
        rows.append([P(r["source"] + "\n" + s.url[:90], small), P(r["type"], small), P(r["status"], small), P(r["confidence"], small), P(r["priority"], small)])
    story.append(table(rows, [8 * cm, 3 * cm, 3.5 * cm, 1.5 * cm, 1.5 * cm]))
    # Inconsistencias
    story.append(PageBreak())
    story.append(P("4. Hallazgos confirmados", h2))
    conf = [s for s in data.sources if s.overall_status == "CONFIRMED_INCONSISTENCY"]
    if not conf:
        story.append(P("No se han confirmado inconsistencias con las evidencias disponibles.", body))
    for s in conf:
        story.append(P(f"{source_row(s)['source']} — {s.url}", small))
        for k, v in (s.field_notes or {}).items():
            if v:
                story.append(P(f"   {k}: {'; '.join(v)}", small))
        for k, v in (s.evidence or {}).items():
            if v:
                story.append(P(f"   Evidencia ({k}): {v}", small))
    story.append(P("5. Hipótesis pendientes de verificación", h2))
    hyp = [s for s in data.sources if s.overall_status in ("POSSIBLE_INCONSISTENCY", "MANUAL_REVIEW", "POSSIBLE_DUPLICATE")]
    if not hyp:
        story.append(P("Sin hipótesis pendientes.", body))
    for s in hyp[:80]:
        story.append(P(f"[{STATUS_LABELS_ES.get(s.overall_status)}] {source_row(s)['source']} — {s.url}: {s.recommended_action or ''}", small))
    for g in data.groups:
        story.append(P(f"[POSIBLE DUPLICADO · {g.status}] {g.platform}: {'; '.join(g.reasons)}. {' '.join(g.warnings)}", small))
    # Plan de corrección
    story.append(P("6. Plan de corrección priorizado", h2))
    rows = [["Prioridad", "Certeza", "Acción", "Detalle"]]
    for x in data.actions[:120]:
        rows.append([P(x.priority, small), P("Confirmada" if x.certainty == "confirmed" else "Hipótesis", small), P(x.title, small), P(x.detail, small)])
    story.append(table(rows, [1.6 * cm, 2 * cm, 6.4 * cm, 7.5 * cm]))
    # Limitaciones
    story.append(P("7. Limitaciones del análisis", h2))
    for lim in (a.limitations or []) + [
        "La comparación depende de que los datos sean públicos y accesibles en el momento de la consulta.",
        "Las clasificaciones automáticas deben revisarse antes de solicitar cambios a terceros.",
    ]:
        story.append(P("• " + lim, body))
    story.append(Spacer(1, 0.5 * cm))
    story.append(P(f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')} por NAP Entity Auditor Pro.", small))
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()
